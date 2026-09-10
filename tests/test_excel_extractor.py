import sys
import types
from collections.abc import Callable
from pathlib import Path
from typing import Any, cast

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet


def _install_dify_plugin_stub() -> None:
    if "dify_plugin" in sys.modules:
        return

    dify_plugin_module = types.ModuleType("dify_plugin")

    class Tool:  # pragma: no cover - minimal import stub
        pass

    setattr(dify_plugin_module, "Tool", Tool)

    entities_module = types.ModuleType("dify_plugin.entities")

    tool_module = types.ModuleType("dify_plugin.entities.tool")

    class ToolInvokeMessage:  # pragma: no cover - minimal import stub
        pass

    setattr(tool_module, "ToolInvokeMessage", ToolInvokeMessage)

    file_package = types.ModuleType("dify_plugin.file")
    file_module = types.ModuleType("dify_plugin.file.file")

    class File:  # pragma: no cover - minimal import stub
        pass

    setattr(file_module, "File", File)
    setattr(file_package, "file", file_module)

    sys.modules["dify_plugin"] = dify_plugin_module
    sys.modules["dify_plugin.entities"] = entities_module
    sys.modules["dify_plugin.entities.tool"] = tool_module
    sys.modules["dify_plugin.file"] = file_package
    sys.modules["dify_plugin.file.file"] = file_module


_install_dify_plugin_stub()

from tools.excel_extractor import ExcelExtractorTool


@pytest.fixture
def excel_tool() -> ExcelExtractorTool:
    return object.__new__(ExcelExtractorTool)


@pytest.fixture
def excel_file_factory(
    excel_tool: ExcelExtractorTool, tmp_path: Path
) -> Callable[..., Any]:
    """Build a real .xlsx on disk and return a mocked File wrapper factory."""

    def _create(
        filename: str, row_text: Callable[[int], str], rows: int = 100
    ) -> Any:
        from dify_plugin.file.file import File

        workbook = Workbook()
        sheet = cast(Worksheet, workbook.active)
        sheet.title = "Sheet1"
        for r in range(rows):
            sheet.append([row_text(r)])
        temp_path = tmp_path / filename
        workbook.save(temp_path)

        fake_file = File()
        fake_file.filename = filename
        fake_file.blob = temp_path.read_bytes()
        fake_file.mime_type = (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

        excel_tool.create_text_message = lambda text: text
        excel_tool.create_blob_message = lambda blob, meta: (blob, meta)
        return fake_file

    return _create


def test_render_row_text_preserves_internal_blanks_and_trims_trailing_blanks(
    excel_tool: ExcelExtractorTool,
) -> None:
    row_text = excel_tool._render_row_text((None, "B", None, "C", None))

    assert row_text == " | B |  | C"


def test_extract_text_xlsx_preserves_blank_cells_between_values(
    excel_tool: ExcelExtractorTool, tmp_path: Path
) -> None:
    workbook = Workbook()
    sheet = cast(Worksheet, workbook.active)
    sheet.title = "Sheet1"
    sheet.append(["A", None, "C"])
    sheet.append(["A", "B", "C"])
    temp_path = tmp_path / "blank-cells.xlsx"
    workbook.save(temp_path)

    extracted_text = excel_tool._extract_text_xlsx(str(temp_path))

    assert "Row 1: A |  | C" in extracted_text
    assert "Row 2: A | B | C" in extracted_text


def test_extract_images_from_wps_etcellimagedata(
    excel_tool: ExcelExtractorTool, monkeypatch: pytest.MonkeyPatch
) -> None:
    import io
    import zipfile
    import olefile

    bio = io.BytesIO()
    with zipfile.ZipFile(bio, "w") as zf:
        zf.writestr("xl/media/image1.png", b"fake_png_data")
        zf.writestr("dummy.txt", b"not an image")
    zip_bytes = bio.getvalue()

    class MockOleFile:
        def __init__(self, *args, **kwargs):
            pass

        def listdir(self, streams=True):
            return [["ETCellImageData"]]

        def openstream(self, entry):
            class MockStream:
                def read(self):
                    return zip_bytes

                def close(self):
                    pass

            return MockStream()

        def close(self):
            pass

    monkeypatch.setattr(olefile, "isOleFile", lambda *args, **kwargs: True)
    monkeypatch.setattr(olefile, "OleFileIO", MockOleFile)

    images = list(excel_tool._extract_images_from_ole_streams("dummy_path.xls"))
    assert len(images) == 1
    assert images[0] == (b"fake_png_data", ".png")


def test_invoke_max_characters_truncation(
    excel_tool: ExcelExtractorTool, excel_file_factory: Callable[..., Any]
) -> None:
    fake_file = excel_file_factory(
        "long.xlsx",
        lambda r: f"Data value {r} with some extra long text content to exceed limit",
    )

    # 1. No max_characters / empty input: full text returned (unlimited)
    for empty_val in [{}, {"max_characters": None}, {"max_characters": ""}, {"max_characters": "   "}]:
        payload = {"excel_content": fake_file, **empty_val}
        messages_unlimited = list(excel_tool._invoke(payload))
        full_text = messages_unlimited[0]
        assert len(full_text) > 500
        assert "[内容已截断" not in full_text

    # 2. With max_characters=200 (int and numeric string): plain cut to 200 chars.
    # No notice suffix is appended: the result must be a verbatim prefix of
    # the unlimited text, so downstream content is never polluted.
    # String parsing itself is covered by test_parse_max_characters; here we
    # just keep one _invoke pass to prove the string value flows through.
    for limit_value in (200, "200"):
        messages_limited = list(
            excel_tool._invoke(
                {"excel_content": fake_file, "max_characters": limit_value}
            )
        )
        limited_text = messages_limited[0]
        assert len(limited_text) == 200
        assert "[内容已截断" not in limited_text
        assert limited_text == full_text[:200]


def test_invoke_max_rows_caps_rows_per_sheet(
    excel_tool: ExcelExtractorTool, excel_file_factory: Callable[..., Any]
) -> None:
    fake_file = excel_file_factory("many-rows.xlsx", lambda r: f"Row data {r}")

    messages = list(
        excel_tool._invoke({"excel_content": fake_file, "max_rows": 10})
    )
    text = messages[0]
    assert "Row 10:" in text
    assert "Row 11:" not in text


def test_extract_text_xlsx_max_columns(
    excel_tool: ExcelExtractorTool, tmp_path: Path
) -> None:
    workbook = Workbook()
    sheet = cast(Worksheet, workbook.active)
    sheet.title = "Sheet1"
    sheet.append(["A", "B", "C"])
    sheet.append(["A", "B", "C"])
    temp_path = tmp_path / "wide.xlsx"
    workbook.save(temp_path)

    extracted_text = excel_tool._extract_text_xlsx(str(temp_path), max_columns=2)

    assert "Row 1: A | B" in extracted_text
    assert "| C" not in extracted_text


def test_parse_positive_int(
    excel_tool: ExcelExtractorTool,
) -> None:
    assert excel_tool._parse_positive_int(None) is None
    assert excel_tool._parse_positive_int("") is None
    assert excel_tool._parse_positive_int(0) is None
    assert excel_tool._parse_positive_int(-3) is None
    assert excel_tool._parse_positive_int(True) is None
    assert excel_tool._parse_positive_int(10) == 10
    assert excel_tool._parse_positive_int("10") == 10


def test_truncate_text_within_limit_returns_unchanged(
    excel_tool: ExcelExtractorTool,
) -> None:
    text = "x" * 500
    assert excel_tool._truncate_text(text, 1000) == text
    assert excel_tool._truncate_text(text, 500) == text


def test_truncate_text_small_limit_never_exceeds_limit(
    excel_tool: ExcelExtractorTool,
) -> None:
    text = "x" * 500
    for limit in (1, 5, 10):
        truncated = excel_tool._truncate_text(text, limit)
        assert len(truncated) <= limit


def test_parse_max_characters(
    excel_tool: ExcelExtractorTool,
) -> None:
    assert excel_tool._parse_max_characters(None) is None
    assert excel_tool._parse_max_characters("") is None
    assert excel_tool._parse_max_characters("   ") is None
    assert excel_tool._parse_max_characters(0) is None
    assert excel_tool._parse_max_characters(-5) is None
    assert excel_tool._parse_max_characters("abc") is None
    assert excel_tool._parse_max_characters(True) is None
    assert excel_tool._parse_max_characters(False) is None
    assert excel_tool._parse_max_characters(float("inf")) is None
    assert excel_tool._parse_max_characters(float("nan")) is None
    assert excel_tool._parse_max_characters(200) == 200
    assert excel_tool._parse_max_characters("200") == 200
    assert excel_tool._parse_max_characters("200.7") == 200
    assert excel_tool._parse_max_characters(200.7) == 200
